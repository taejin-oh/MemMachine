"""memmachine_workload.pptx(3장)에 이어지는 추가 7장(4~10장)을 만든다.

구성 의도
  4장  배경      — 세션·인스턴스·캐시가 무엇이고 요청 1건이 겪는 세 상황
  5장  선택      — 사용자를 나누는 두 방식(세션 분리 / 이름표 분리)과 현재 상태
  6장  실측      — 상태별 저장소 질의 배수
  7장  실측 상세 — 신규 등록 18회의 내역
  8장  실측 상세 — worker를 늘렸을 때
  9장  구조      — 컬렉션 둘, 세션은 payload 한 칸
  10장 인접 비용 — 임베딩 호출 분할

기준 커밋 a8322a7. 수치는 docs/msr/qdrant_requests/ 의 보고서와 상세 문서에서만 가져왔다.
시각 체계는 deck_style.py가 원본 덱에서 추출한 토큰 그대로다.
"""

from pptx.enum.text import PP_ALIGN

import deck_style as S

OUT = "memmachine_qdrant_requests.pptx"
M = 0.45          # 좌우 여백
FULL = 12.43      # 전폭 콘텐츠


def _labeled_box(sl, x, y, w, h, kind, title, lines, *, badge=None,
                 title_size=11, line_size=9, gap=0.20, dashed=False):
    """색 박스 + 제목 + 본문 줄. box()는 가운데 정렬이라 직접 얹는다."""
    S.box(sl, x, y, w, h, [""], kind=kind, dashed=dashed)
    if badge:
        S.badge(sl, x + w - 0.30, y - 0.06, badge)
    S.text(sl, x + 0.15, y + 0.08, w - 0.45, 0.22, title, size=title_size, bold=True)
    for i, line in enumerate(lines):
        S.text(sl, x + 0.15, y + 0.32 + i * gap, w - 0.25, 0.18, line, size=line_size)


# ---------------------------------------------------------------- 4장 배경

def slide4(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "먼저 용어 셋 — 세션, 인스턴스, 캐시",
        "이 세 가지를 구분해야 뒤의 숫자가 읽힌다  ·  기준 upstream/speedkick a8322a7",
    )

    S.section(sl, M, 1.14, "은행 창구에 비유하면", w=6.0)

    for i, (kind, title, lines) in enumerate([
        ("store", "세션 = 고객의 계좌",
         ["기억을 담는 칸. 고객사와 프로젝트 조합으로 하나씩 만들어진다",
          "내용은 데이터베이스에 영구 보관된다"]),
        ("self", "인스턴스 = 창구 직원의 업무 세팅",
         ["그 계좌를 다루려고 서버 메모리에 올려둔 준비물",
          "금고 번호를 적은 메모와 열쇠이지, 돈이 아니다"]),
        ("opt", "캐시 = 책상 위 세팅 100칸",
         ["자주 쓰는 준비물만 올려둔다",
          "10분 안 쓰면 치우고, 자리가 모자라면 오래된 것부터 치운다"]),
    ]):
        _labeled_box(sl, M, 1.48 + i * 0.78, 6.05, 0.70, kind, title, lines,
                     line_size=8.5)

    S.rect(sl, M, 3.86, 6.05, 0.92, S.WHITE, S.API_L, 1.5)
    S.text(sl, M + 0.15, 3.94, 5.7, 0.22, "인스턴스에는 기억이 들어 있지 않다",
           size=11, bold=True, color=S.API_L)
    S.text(sl, M + 0.15, 4.18, 5.85, 0.40,
           "한 세션에 기억이 수백만 건 쌓여도 인스턴스 크기는 거의 변하지 않는다. "
           "그래서 100칸을 채워도 메모리를 크게 쓰지 않는다.", size=9, wrap=True)

    S.text(sl, 6.60, 1.14, 6.28, 0.28, "인스턴스 안에 무엇이 있나",
           size=14, bold=True, color=S.NAVY)
    S.table(
        sl, 6.60, 1.48, 6.28, [3.14, 3.14],
        [
            ["저장소의 어느 칸인지 가리키는 주소", "기억 본문"],
            ["데이터베이스 연결 통로", "기억을 숫자로 바꾼 값"],
            ["쪼개고 변환하는 방식 설정", "검색 결과"],
        ],
        header=["들어 있는 것 (도구)", "없는 것 (항상 저장소에서)"],
        header_size=9, body_size=8.5, row_h=0.28,
    )

    S.text(sl, 6.60, 3.10, 6.28, 0.28, "요청 1건이 겪는 세 가지 상황",
           size=14, bold=True, color=S.NAVY)
    S.table(
        sl, 6.60, 3.44, 6.28, [2.38, 1.20, 1.10, 1.60],
        [
            ["바로 처리 — 세팅이 책상에 있다", "0", "1", "1회"],
            ["재준비 — 계좌는 있고 세팅만 없다", "1", "1", "2회"],
            ["신규 등록 — 계좌 자체가 없다", "17", "1", "18회"],
        ],
        header=["상황", "준비 질의", "본래 일", "합계"],
        header_size=9, body_size=8.5, row_h=0.30,
        mono_cols={1, 2, 3}, emphasis_rows={2},
    )

    S.conclusion(sl, M, 5.15, FULL, 1.55, [
        ("계좌는 영구히 남지만 업무 세팅은 수시로 치워진다.", " 치워졌다고 기억이 사라지는 것은 아니다."),
        ("인스턴스가 있어도 검색은 반드시 저장소에 한 번 간다.", " 캐시가 아끼는 것은 검색이 아니라 그 앞의 준비 동작이다."),
        ("따라서 요청 1건의 비용은", " 요청 종류가 아니라 그 순간의 준비 상태로 정해진다."),
        ("이 세 상황의 이름을 뒤에서 계속 쓴다.", " 바로 처리, 재준비, 신규 등록."),
    ])
    S.footer(sl, "근거: 보고서 2부(이 문서에서 쓰는 말) — 캐시 용량 100과 유휴 600초는 "
                 "episodic_memory_manager.py:43, :48. 이 파일은 작업 브랜치와 speedkick이 동일하다.",
             y=6.82)


# ---------------------------------------------------------------- 5장 두 방식

def slide5(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "사용자를 나누는 두 가지 방법 — 먼저 고르고 나서 측정해야 한다",
        "두 방식은 서로 다른 병목을 드러낸다  ·  지금 평가 코드는 둘 중 어느 쪽도 아니다",
    )

    S.box(sl, M, 1.14, 6.05, 2.52, [""], kind="store")
    S.text(sl, M + 0.15, 1.22, 5.7, 0.24, "A. 세션 분리 — 사용자마다 방 하나",
           size=12, bold=True)
    S.text(sl, M + 0.15, 1.50, 5.85, 0.18,
           "사용자 한 명에게 계좌 하나. 저장소 안에서 칸이 나뉜다", size=9)
    S.rect(sl, M + 0.15, 1.74, 5.75, 0.30, S.WHITE, S.STORE_L, 1.0)
    S.text(sl, M + 0.24, 1.80, 5.6, 0.18,
           "sys-partition_key = 사용자마다 다른 값", size=8.5, font=S.MONO, color=S.CODE)
    S.text(sl, M + 0.15, 2.14, 5.7, 0.2, "좋은 점", size=10, bold=True, color=S.STORE_L)
    for i, line in enumerate([
        "자기 방만 보는 것이 기본 동작이다. 조건을 빠뜨려도 안전하다",
        "사용자 한 명을 지우려면 방을 통째로 지우면 된다",
    ]):
        S.text(sl, M + 0.15, 2.36 + i * 0.19, 5.85, 0.18, line, size=8.5)
    S.text(sl, M + 0.15, 2.82, 5.7, 0.2, "비용", size=10, bold=True, color=S.API_L)
    for i, line in enumerate([
        "사용자마다 신규 등록 17회. 500명이면 8,500회",
        "동시 활동 사용자가 100명을 넘으면 재준비가 반복된다",
        "데이터베이스에도 사용자 수만큼 칸이 생긴다",
    ]):
        S.text(sl, M + 0.15, 3.04 + i * 0.19, 5.85, 0.18, line, size=8.5)

    S.box(sl, 6.83, 1.14, 6.05, 2.52, [""], kind="opt")
    S.text(sl, 6.98, 1.22, 5.7, 0.24, "B. 이름표 분리 — 한 방을 같이 쓰기",
           size=12, bold=True)
    S.text(sl, 6.98, 1.50, 5.85, 0.18,
           "모두 한 계좌를 쓰고, 기억마다 누구 것인지 표시한다", size=9)
    S.rect(sl, 6.98, 1.74, 5.75, 0.30, S.WHITE, S.OPT_L, 1.0)
    S.text(sl, 7.07, 1.80, 5.6, 0.18,
           "producer_id = 'userA' OR produced_for_id = 'userA'",
           size=8.5, font=S.MONO, color=S.CODE)
    S.text(sl, 6.98, 2.14, 5.7, 0.2, "좋은 점", size=10, bold=True, color=S.STORE_L)
    for i, line in enumerate([
        "방을 만드는 등록 절차가 전체를 통틀어 한 번이다",
        "책상도 한 칸만 쓴다. 100칸 한도 문제가 생기지 않는다",
    ]):
        S.text(sl, 6.98, 2.36 + i * 0.19, 5.85, 0.18, line, size=8.5)
    S.text(sl, 6.98, 2.82, 5.7, 0.2, "비용과 위험", size=10, bold=True, color=S.API_L)
    for i, line in enumerate([
        "검색마다 조건을 붙여야 한다. 빠뜨리면 남의 기억이 나온다",
        "말한 사람만으로 걸면 상대 답변이 빠진다. 두 조건이 필요하다",
        "사용자 한 명만 지우려면 기억을 하나씩 찾아 지워야 한다",
    ]):
        S.text(sl, 6.98, 3.04 + i * 0.19, 5.85, 0.18, line, size=8.5)

    S.table(
        sl, M, 3.78, FULL, [4.00, 4.20, 4.23],
        [
            ["신규 등록 비용", "사용자 수 × 17회", "전체 1회"],
            ["캐시 100칸 한도", "100명 넘으면 재준비 반복", "문제 없음"],
            ["격리 방식", "구조적 — 조건을 빠뜨려도 안전", "정책적 — 조건을 붙여야 안전"],
        ],
        header=["갈리는 지점", "A. 세션 분리", "B. 이름표 분리"],
        header_size=9, body_size=9, row_h=0.26,
    )

    S.rect(sl, M, 4.90, FULL, 0.60, S.WHITE, S.API_L, 1.5)
    S.badge(sl, M + FULL - 0.32, 4.84, "yes")
    S.text(sl, M + 0.15, 4.97, 7.0, 0.22, "지금 평가 코드는 둘 중 어느 쪽도 아니다",
           size=11, bold=True, color=S.API_L)
    S.text(sl, M + 0.15, 5.20, 12.0, 0.2,
           "방을 나누지 않아 모든 사용자가 기본 세션 하나로 합쳐지고, 검색에 조건을 걸지 않는다. "
           "이름표(화자)는 붙지만 격리에 쓰이지 않는다", size=9)

    S.conclusion(sl, M, 5.58, FULL, 1.28, [
        ("둘 중 하나가 정답인 것은 아니다.", " 격리가 중요하면 A, 사용자 수가 많고 운영 비용이 중요하면 B다."),
        ("다만 먼저 고르지 않으면 측정 설계가 서지 않는다.", " A는 칸이 늘 때를, B는 한 칸이 커질 때를 본다."),
        ("현재 상태로는 사용자를 500명까지 늘려도", " 두 방식의 비용 차이가 드러나지 않는다."),
    ], gap=0.25, size=10)
    S.footer(sl, "근거: 보고서 7부(두 방식)·8부(현재 평가) — 세션 키는 service.py:44, 기본값 universal은 spec.py:29. "
                 "필터 회수율은 사용자 3명 12건 적재로 실측했다.", y=6.94)


# ---------------------------------------------------------------- 6장 실측 배수

def slide6(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "요청 1건이 저장소 질의 몇 건이 되는가 — 상태별 실측",
        "기준 acb4f9a → a8322a7 (speedkick, 2026-09-14) · qdrant 1.17.0 + postgres 16 · HTTP 전수 기록",
    )

    cols = [3.75, 0.95, 0.95, 0.95, 0.95]
    bottom = S.table(
        sl, M, 1.20, 7.55, cols,
        [
            ["바로 처리: 인스턴스가 캐시에 있음", "1", "1", "1", "3"],
            ["재준비: 캐시에서 빠짐, 세션은 등록됨", "2", "2", "2", "4"],
            ["신규 등록: 처음 보는 세션", "18", "18", "—", "—"],
        ],
        header=["상황", "저장 1건", "검색 1건", "에피소드 삭제", "세션 삭제"],
        header_size=9, body_size=9, mono_cols={1, 2, 3, 4}, emphasis_rows={2},
    )
    for i in range(3):
        S.badge(sl, 0.16, 1.51 + i * 0.30, "yes")

    S.rect(sl, M, bottom, 7.55, 0.30, S.SELF_F, S.WHITE, 1.0)
    S.text(sl, M + 0.07, bottom + 0.05, 3.61, 0.2, "프로세스(worker) 시작", size=9)
    S.text(sl, M + 3.82, bottom + 0.05, 3.66, 0.2,
           "연결 확인 1회 + 버전 확인 1회 (별도 스레드)", size=9)
    S.text(sl, M, bottom + 0.34, 7.55, 0.2,
           "GET /collections는 트레이스로, GET /는 클라이언트 버전 경고로 확인", size=8, color=S.FOOT)

    _labeled_box(sl, M, 3.15, 7.55, 1.30, "opt",
                 "설정에 따라서만 배수가 달라지는 셋 — 모두 코드 기준, 실행 미확인",
                 ["agent_mode=true → 검색 1건이 저장소 2~22회",
                  "시맨틱 vector_store 백엔드 → set_id마다 limit 10,000 검색 1회",
                  "분산 모드 → 세션 생성 때 샤드 키 1회 추가"],
                 badge="part", gap=0.26)

    S.section(sl, 8.25, 1.16, "배수를 정하는 것은 두 가지뿐", w=4.63)
    S.box(sl, 8.25, 1.50, 4.63, 0.44,
          ["① 인스턴스가 이 worker의 캐시에 있는가 (100칸 / 10분)"], kind="self", size=9)
    S.box(sl, 8.25, 1.98, 4.63, 0.44,
          ["② 그 세션이 저장소에 등록돼 있는가"], kind="self", size=9)
    S.text(sl, 8.25, 2.46, 4.63, 0.2, "둘 다 A1 / S1(세션 확보) 한 단계에서 갈린다",
           size=9, color=S.GRAY)

    S.box(sl, 8.25, 2.70, 4.63, 0.78, [""], kind="self")
    S.text(sl, 8.40, 2.78, 4.3, 0.22, "검색 옵션은 배수를 바꾸지 않는다", size=11, bold=True)
    S.text(sl, 8.40, 3.00, 4.35, 0.2,
           "top_k · filter · expand_context · score_threshold 모두 1회", size=9)
    S.text(sl, 8.40, 3.20, 4.35, 0.2, "바뀌는 것은 limit = top_k × 4 하나뿐",
           size=8.5, font=S.MONO, color=S.CODE)

    S.box(sl, 8.25, 3.56, 4.63, 0.60, [""], kind="opt")
    S.text(sl, 8.40, 3.63, 4.3, 0.22, "기준 커밋이 바뀌었다", size=11, bold=True)
    S.text(sl, 8.40, 3.85, 4.35, 0.2,
           "acb4f9a (1~3장) → a8322a7 · 캐시 코드는 작업 브랜치와 동일",
           size=8.5, font=S.MONO, color=S.CODE)

    S.panel(sl, 8.25, 4.24, 4.63, 0.86, radius=0.10)
    S.text(sl, 8.40, 4.31, 4.3, 0.2, "배지 뜻이 바뀐다 — 1~3장은 계측 유무, 4장부터는 근거 종류",
           size=9, bold=True)
    for i, (kind, label) in enumerate([("yes", "실행 트레이스로 실측"),
                                       ("part", "코드 기준, 실행 미확인"),
                                       ("no", "이번 회차 계측 밖")]):
        yy = 4.55 + i * 0.17
        S.badge(sl, 8.42, yy - 0.01, kind, d=0.15, size=7)
        S.text(sl, 8.62, yy, 4.1, 0.16, label, size=8.5)

    S.conclusion(sl, M, 5.18, FULL, 1.58, [
        ("평상시 배수는 1이다.", " 저장 1건도 검색 1건도 저장소 질의 1건이고, 검색 옵션은 이 수를 바꾸지 않는다."),
        ("배수를 키우는 것은 검색 옵션이 아니라 준비 상태다.", " 재준비 2배, 신규 등록 18배."),
        ("삭제는 값이 다르다.", " 에피소드 삭제 1회, 세션 삭제는 바로 처리 3회 재준비 4회다."),
        ("따라서 부하 산정은", " 요청 수가 아니라 요청 수 × 상태 분포 × worker 수로 해야 한다."),
    ])
    S.footer(sl, "근거: 보고서 6부(상태별 요청 수) — qdrant-client의 HTTP 전송 함수를 감싸 REST 요청을 전수 기록했고 "
                 "실제 EpisodicMemoryManager와 세션 DB를 그대로 썼다. 임베더만 64차원 해시.", y=6.86)


# ---------------------------------------------------------------- 7장 신규 등록

def slide7(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "신규 등록 18회의 내역 — 인덱스 11회는 낭비가 아니라 의도된 동작",
        "PUT /collections ×2 + PUT …/index ×11 + registry 왕복 ×4 = 17  (+ 본 요청 1 = 18)",
    )

    cx = M + 5.45 / 2 - 0.055
    S.box(sl, M, 1.20, 5.45, 0.58,
          [("1~3  registry 확인과 생성", True), "포인트 조회 → registry 컬렉션 PUT → 포인트 재조회"],
          kind="self", size=9)
    S.arrow_down(sl, cx, 1.81)
    S.box(sl, M, 1.99, 5.45, 0.50,
          [("4  native 컬렉션 생성 PUT", True), "hnsw m=0, payload_m=16"], kind="store", size=9)
    S.arrow_down(sl, cx, 2.52)

    S.rect(sl, M, 2.70, 5.45, 1.72, S.WHITE, S.API_L, 2.0)
    S.text(sl, M + 0.15, 2.78, 5.15, 0.22,
           "5~15  payload 인덱스 생성 11회 — 이미 있어도 매번 나간다",
           size=10.5, bold=True, color=S.API_L)
    idx = ["sys-partition_key (is_tenant)", "_timestamp (datetime)", "_episode_uid",
           "_session_key", "_producer_id", "_producer_role", "_produced_for_id",
           "_sequence_num (integer)", "_episode_type", "_content_type",
           "_created_at (datetime)"]
    for i, name in enumerate(idx):
        col, row = divmod(i, 6)
        S.text(sl, M + 0.17 + col * 2.65, 3.06 + row * 0.21, 2.55, 0.18, name,
               size=8, font=S.MONO, color=S.CODE)

    S.arrow_down(sl, cx, 4.45)
    S.box(sl, M, 4.63, 5.45, 0.46, ["16~17  registry 포인트 upsert → 최종 조회"],
          kind="self", size=9)

    S.table(
        sl, 6.20, 1.20, 6.68, [3.28, 1.70, 1.70],
        [
            ["1  registry 포인트 조회", "404", "200, 빈 결과"],
            ["2  registry 컬렉션 생성", "200", "409"],
            ["3  registry 포인트 조회", "200, 빈 결과", "200, 빈 결과"],
            ["4  native 컬렉션 생성", "200", "409"],
            ["5~15  payload 인덱스 11회", "200", "200"],
            ["16  registry 포인트 등록", "200", "200"],
            ["17  registry 포인트 조회", "200", "200"],
            ["합계", "17회", "17회"],
        ],
        header=["요청", "첫 세션", "두 번째 새 세션"],
        header_size=9, body_size=8.5, row_h=0.28, center_cols={1, 2},
        emphasis_rows={7}, emphasis_color=S.NAVY,
    )
    S.badge(sl, 12.52, 3.46, "yes")

    _labeled_box(sl, 6.20, 3.88, 6.68, 1.05, "store",
                 "인덱스 11회를 매번 보내는 이유",
                 ["예전에는 컬렉션이 이미 있으면 payload 인덱스가 하나도 만들어지지 않았다",
                  "그 결함을 고치면서 컬렉션 생성과 인덱스 생성을 분리해 개별로 보낸다",
                  "커밋 a2a1754 · 첫 세션 각 32~88ms, 두 번째부터 각 3ms (로컬 Docker)"],
                 gap=0.21)

    S.conclusion(sl, M, 5.15, FULL, 1.55, [
        ("17회 중 11회는 인덱스 생성이지만 낭비가 아니다.", " 인덱스가 하나도 만들어지지 않던 결함을 막기 위한 의도된 동작이다."),
        ("두 번째 새 세션부터는", " 컬렉션 생성 2회가 409로 돌아온다. 총 요청 수는 17회로 같다."),
        ("검색도 같은 경로를 쓴다.", " 처음 보는 세션으로 검색만 해도 17회 + 검색 1회가 나간다."),
        ("방식 A를 고르면", " 이 17회가 사용자 수만큼 그대로 누적된다."),
    ])
    S.footer(sl, "근거: 보고서 4.3~4.4(요청 상세) · 트레이스 원본 trace_output.txt — 코드는 a8322a7 기준 "
                 "service_locator.py:106·118·123, qdrant_vector_store.py:649-704.", y=6.82)


# ---------------------------------------------------------------- 8장 worker

def slide8(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "worker를 늘리면 요청은 몇 배가 되는가 — 캐시는 worker마다 따로 있다",
        "MemoryInstanceCache  capacity=100  idle=600s  sweep=2s  ·  MEMMACHINE_WORKERS 기본값 1",
    )

    S.table(
        sl, M, 1.20, 6.15, [1.75, 2.20, 2.20],
        [
            ["Qdrant 클라이언트", "프로세스 수명, 시작 시 예열", "GET /collections 1회"],
            ["인스턴스 캐시", "worker별, 100칸 · 유휴 600초", "조회 1회, 없으면 생성 17회"],
            ["registry 엔트리", "캐시하지 않는다", "열 때마다 조회 1회"],
            ["시맨틱 컬렉션 핸들", "서비스 시작 시 1회", "있으면 1회, 없으면 16회"],
        ],
        header=["캐시 층", "범위와 수명", "미스일 때 저장소 질의"],
        header_size=9, body_size=8.5, row_h=0.32,
        emphasis_rows={1}, emphasis_color=S.NAVY,
    )
    S.badge(sl, 0.16, 1.51, "yes")
    S.badge(sl, 0.16, 2.47, "part")

    _labeled_box(sl, M, 2.92, 6.15, 0.75, "self",
                 "캐시에서 빠지는 경우는 셋뿐",
                 ["worker 재시작 · 10분 넘게 쉰 세션 · worker당 활성 세션 100개 초과",
                  "100칸 · 600초 · 2초는 설정으로 바꿀 수 없다"],
                 gap=0.24)

    S.table(
        sl, 6.78, 1.20, 6.10, [2.90, 1.30, 1.90],
        [
            ["worker N개 동시 재시작 직후", "2 × N", "연결 + 버전 확인"],
            ["기존 세션이 각 worker에 처음 닿음", "+1 × N", "조회가 worker마다"],
            ["새 세션이 M개 생김", "17 × M", "worker 수와 무관"],
            ["같은 새 세션을 두 worker가 동시에", "17 × 2", "잠금이 프로세스 안"],
            ["worker당 활성 세션 100개 초과", "모두 2배", "캐시 용량 초과"],
            ["10분 유휴 뒤 첫 요청", "세션당 +1", "유휴 600초 제거"],
        ],
        header=["운영 시나리오", "저장소 질의", "근거"],
        header_size=9, body_size=8.5, row_h=0.30, mono_cols={1}, emphasis_rows={2},
    )

    _labeled_box(sl, M, 4.00, FULL, 0.75, "opt",
                 "로드밸런서는 있지만 사용자를 같은 worker로 보내는 설정은 없다",
                 ["그래서 같은 세션이 여러 worker에 중복으로 올라가고, 처음 받는 worker마다 재준비가 한 번씩 발생한다. "
                  "기본값이 worker 1개라 이 문제는 2개 이상으로 올렸을 때만 생긴다"],
                 gap=0.22)

    S.conclusion(sl, M, 4.95, FULL, 1.72, [
        ("캐시는 worker마다 따로 있다.", " worker가 N개면 같은 세션이라도 worker마다 각각 재준비 비용을 낸다."),
        ("다만 신규 등록 17회는 worker 수가 아니라 새 세션 수에 비례한다.", " 먼저 연 worker가 등록하면 나머지는 조회 1회로 끝난다."),
        ("활성 세션이 worker당 100개를 넘으면", " 모든 요청이 재준비가 되어 질의가 계속 2배로 유지된다."),
        ("부하 시험은 worker 수와 세션 수를 함께 변수로 잡아야 한다.", " 하나만 바꾸면 이 배수가 드러나지 않는다."),
    ])
    S.footer(sl, "근거: 보고서 10부(캐시 리뷰) — episodic_memory_manager.py:43-52·103-106, "
                 "instance_lru_cache.py:174-185·206-219, docker-compose.yml:87. "
                 "캐시 잠금 자체는 병목이 아니다(요청당 마이크로초).", y=6.80)


# ---------------------------------------------------------------- 9장 구조

def slide9(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "컬렉션은 둘, 세션은 payload 한 칸 — 세션이 늘어도 컬렉션은 늘지 않는다",
        "long_term_memory__registry  ·  long_term_memory__<sha256(config)>",
    )

    S.box(sl, M, 1.20, 5.75, 1.05, [""], kind="store")
    S.text(sl, M + 0.15, 1.28, 5.4, 0.22, "registry 컬렉션", size=11, bold=True)
    S.text(sl, M + 0.15, 1.50, 5.4, 0.2, "long_term_memory__registry",
           size=9.5, font=S.MONO, color=S.CODE)
    S.text(sl, M + 0.15, 1.70, 5.45, 0.2,
           "세션 하나당 포인트 1개. 벡터는 1차원 더미, payload에 이름과 스키마", size=9)
    S.text(sl, M + 0.15, 1.90, 5.45, 0.2,
           "포인트 id = uuid5(네임스페이스, 파티션 키) — 계산해 바로 조회", size=9)

    S.box(sl, M, 2.35, 5.75, 1.60, [""], kind="store")
    S.text(sl, M + 0.15, 2.43, 5.4, 0.22, "native 컬렉션", size=11, bold=True)
    S.text(sl, M + 0.15, 2.65, 5.4, 0.2, "long_term_memory__<sha256(config)>",
           size=9.5, font=S.MONO, color=S.CODE)
    S.text(sl, M + 0.15, 2.85, 5.45, 0.2,
           "이름이 설정 해시라서 같은 설정을 쓰는 모든 세션이 하나를 공유한다", size=9)
    for i, val in enumerate(["sys-partition_key=bd36…", "…=a1f2…", "…=7c04…"]):
        x = M + 0.17 + i * 1.85
        S.rect(sl, x, 3.10, 1.75, 0.22, S.WHITE, S.STORE_L, 1.0)
        S.text(sl, x + 0.06, 3.13, 1.65, 0.16, val, size=7.5, font=S.MONO, color=S.CODE)
    S.text(sl, M + 0.15, 3.42, 5.45, 0.2,
           "포인트에는 벡터와 payload 11키 + 사용자 metadata만 들어간다", size=9, bold=True)
    S.text(sl, M + 0.15, 3.62, 5.45, 0.2, "본문 텍스트는 저장소에 저장하지 않는다", size=9)

    _labeled_box(sl, M, 4.05, 5.75, 0.88, "self", "파티션 키 규칙",
                 ["세션 키가 규칙에 맞으면 그대로, 아니면 sha256[:32]",
                  "orga/prja → bd36dbe5eef78ec754963d978108ef7b"], gap=0.20)

    S.box(sl, 6.45, 1.20, 6.43, 1.30, [""], kind="store")
    S.badge(sl, 12.58, 1.14, "yes")
    S.text(sl, 6.60, 1.28, 5.6, 0.22, "검색에 실제로 나가는 요청", size=10.5, bold=True)
    for i, line in enumerate([
        "POST /collections/long_term_memory__<hash>/points/query/batch",
        "searches=1  limit=80  with_payload=false  with_vector=false",
        'filter={"must":[{"key":"sys-partition_key","match":{"value":"bd36…"}}]}',
    ]):
        S.text(sl, 6.60, 1.52 + i * 0.19, 6.15, 0.18, line, size=8.5,
               font=S.MONO, color=S.CODE)
    S.text(sl, 6.60, 2.13, 6.15, 0.2,
           "방식 B의 사용자 조건은 이 파티션 조건과 AND로 묶인다", size=9)

    S.rect(sl, 6.45, 2.60, 6.43, 1.35, S.WHITE, S.API_L, 1.5)
    S.badge(sl, 12.58, 2.54, "no")
    S.text(sl, 6.60, 2.68, 5.6, 0.22, "한 칸에 약 1,667건 전까지는 그래프가 없다",
           size=11.5, bold=True, color=S.API_L)
    S.text(sl, 6.60, 2.94, 6.15, 0.60,
           "m=0이라 전역 그래프는 없고, payload_m=16으로 파티션 값마다 하위 그래프가 생긴다. "
           "임계값은 1536차원 약 1,667개다. 방식 A는 사용자당, 방식 B는 프로젝트 전체가 기준이 된다.",
           size=9, wrap=True)
    S.text(sl, 6.60, 3.66, 6.15, 0.18,
           "이전 Qdrant 서버 소스 분석(v1.19.1) 인용 — 이번 회차에 재계측하지 않았다",
           size=8, color=S.FOOT)

    _labeled_box(sl, 6.45, 4.05, 6.43, 0.88, "self", "탐색 폭과 점수 컷",
                 ["ef 기본 100, 실제 폭은 max(ef, limit) — top_k 25 초과부터 커진다",
                  "score_threshold는 요청에 실리지 않는다 (long_term_memory.py:340)"], gap=0.20)

    S.conclusion(sl, M, 5.10, FULL, 1.58, [
        ("세션이 늘어도 컬렉션은 늘지 않는다.", " 늘어나는 것은 registry 포인트 1개와 파티션 값 하나다."),
        ("세션 분리는 컬렉션이 아니라 payload 필드 하나로 한다.", " sys-partition_key에 걸린 테넌트 인덱스가 그 경계다."),
        ("방식 A는 사용자 대부분이 임계값에 못 미쳐", " 그래프 없이 필터 후 순차 비교로 검색된다. 소규모에서는 이 편이 빠르다."),
        ("방식 B는 그래프가 일찍 생기지만", " 사용자 조건의 선택도에 따라 탐색 경로가 갈린다."),
    ], size=10)
    S.footer(sl, "근거: 보고서 3.1(구조)·4.6(검색 요청)·7.3(임계값) — qdrant_vector_store.py:64·87·550·669-672, "
                 "service_locator.py:158-182, long_term_memory.py:310·340.", y=6.80)


# ---------------------------------------------------------------- 10장 임베딩

def slide10(prs):
    sl = S.add_slide(prs)
    S.header(
        sl,
        "저장이 느려지면 저장소가 아닐 수 있다 — 임베딩 호출의 분할",
        "한 요청 상한 inputs ≤ 2048, chars ≤ 75000  ·  max_attempts=1  ·  재시도 없음",
    )

    S.box(sl, M, 1.20, FULL, 0.95, [""], kind="api")
    S.text(sl, M + 0.15, 1.27, 4.0, 0.22, "분할 규칙", size=11, bold=True)
    for i, line in enumerate([
        "① batch_size가 설정되면 그 크기로 먼저 나눠 동시에 부른다 (기본 미설정)",
        "② 입력 하나가 75,000자를 넘으면 그 길이로 자른다",
        "③ 한 요청은 입력 2,048개 이하, 총 75,000자 이하로 묶는다",
        "④ 묶음마다 HTTP 요청을 만들어 동시에 보낸다",
    ]):
        col, row = divmod(i, 2)
        S.text(sl, M + 0.15 + col * 6.15, 1.52 + row * 0.20, 6.0, 0.18, line, size=8.5)
    S.text(sl, M + 0.15, 1.93, 12.1, 0.18,
           "조각은 길이 가중 평균으로 합쳐 벡터 1개가 되므로 저장소 포인트 수는 변하지 않는다",
           size=9, color=S.GRAY)

    S.table(
        sl, M, 2.30, 7.15, [3.25, 1.70, 1.00, 1.20],
        [
            ["저장, 에피소드 4건", "4개 336자", "1", "1"],
            ["검색 1건", "1개 16자", "1", "1"],
            ["저장, 16만 자 1건", "1개 160,039자", "1", "3"],
            ["저장, 에피소드 2,100건", "2,100개 186,780자", "1", "3"],
            ["저장, batch_size=2로 5건", "5개", "1", "3"],
            ["저장, sentence_text 2건×3문장", "6개", "1", "1"],
            ["임베더 최초 생성 (프로세스당 1회)", "1개", "1", "1"],
        ],
        header=["구간", "입력", "논리 호출", "HTTP 요청"],
        header_size=9, body_size=8.5, row_h=0.28, mono_cols={2, 3}, center_cols={1},
    )
    S.badge(sl, 0.16, 2.61, "yes")
    S.text(sl, M, 4.58, 7.15, 0.18,
           "HTTP 3건인 경우의 분할: 75,000+75,000+10,039자 · 854+836+410개 · 2+2+1개",
           size=8, color=S.FOOT)

    _labeled_box(sl, 7.75, 2.30, 5.13, 0.95, "opt", "운영값으로 환산하면",
                 ["메시지 평균 400자일 때 한 저장 요청에 약 187건이 넘어가면",
                  "HTTP 요청이 하나 더 생긴다. 2,048개 제한은 짧을 때만 먼저 걸린다"],
                 gap=0.21)

    S.box(sl, 7.75, 3.40, 5.13, 1.15, [""], kind="self")
    S.text(sl, 7.90, 3.48, 4.8, 0.22, "바로 처리 상태의 외부 호출", size=11, bold=True)
    S.text(sl, 7.90, 3.72, 4.85, 0.18, "검색 = 임베딩 1 + 저장소 1 + PostgreSQL 3",
           size=9, font=S.MONO, color=S.CODE)
    S.text(sl, 7.90, 3.91, 4.85, 0.18, "저장 = 임베딩 1+ · 저장소 1 · PostgreSQL 다수",
           size=9, font=S.MONO, color=S.CODE)
    S.text(sl, 7.90, 4.14, 4.85, 0.35,
           "임베딩과 저장소 검색은 직렬이라 임베딩 왕복이 검색 지연에 그대로 더해진다",
           size=9, wrap=True)

    S.rect(sl, M, 4.85, FULL, 0.60, S.WHITE, S.API_L, 1.5)
    S.text(sl, M + 0.15, 4.92, 6.0, 0.22, "임베딩 호출에는 재시도가 없다",
           size=11, bold=True, color=S.API_L)
    S.text(sl, M + 0.15, 5.15, 12.1, 0.2,
           "모든 호출부가 max_attempts 기본값 1을 쓴다. 429나 5xx를 받으면 재시도 없이 저장 또는 검색이 그대로 실패한다. "
           "임베더 안의 백오프 루프는 한 번도 돌지 않는다", size=9)

    S.conclusion(sl, M, 5.60, FULL, 1.35, [
        ("저장소 질의가 1회로 고정이어도 임베딩 HTTP는 1건 이상일 수 있다.", " 저장 요청이 커질 때 갈라지는 쪽은 임베딩이다."),
        ("병목 측정에서 임베딩 왕복과 저장소 왕복을 나눠 재야 한다.", " 검색에서 둘은 직렬이다."),
        ("임베딩 호출 수는 준비 상태와 무관하다.", " 신규 등록이어도 저장 1회, 검색 1회다."),
        ("주기적으로 도는 임베딩은 시맨틱 특징 임베딩뿐이고,", " 기본값은 꺼짐이다."),
    ], size=10, gap=0.25)
    S.footer(sl, "근거: 보고서 9부(임베딩) · 측정 기록 3부(E1~E11) — OpenAI 호환 스텁 서버로 실제 HTTP 요청의 "
                 "입력 개수와 총 길이를 셌고, 임베더 구현은 운영과 같은 OpenAIEmbedder다.", y=7.06)


def main():
    prs = S.new_deck()
    for fn in (slide4, slide5, slide6, slide7, slide8, slide9, slide10):
        fn(prs)
    prs.save(OUT)
    print(f"saved {OUT}: {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    main()
