"""memmachine_workload.pptx의 시각 체계를 그대로 재현하는 헬퍼.

원본에서 추출한 토큰만 사용한다.
- 슬라이드 13.33x7.5in, Blank 레이아웃에 도형 직접 배치
- 박스: 둥근사각형 adj=0.12, 선 1.3pt, 글자 가운데/중앙 정렬, 좌여백 0.04in
- 배지: 원형 0.24in, 흰 테두리 1.2pt, 9pt 굵은 흰 글자
- 결론 패널 FAFAFA/999999, 정보 패널 F3F6FA/8CA0BE
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---- 글꼴 ----
KR = "Apple SD Gothic Neo"
MONO = "Menlo"

# ---- 색 ----
NAVY = "1F355E"
WHITE = "FFFFFF"
SUBTITLE = "C8D4E6"
BODY = "1A1A1A"
GRAY = "555555"
FOOT = "888888"
CODE = "1F4E79"
RED_TXT = "A33333"

SELF_F, SELF_L = "DBE6F6", "2F4A7A"      # MemMachine 자체
STORE_F, STORE_L = "CDEAE4", "1F7A6A"    # 외부 저장소
API_F, API_L = "F9DCDC", "A33333"        # 외부 API
OPT_F, OPT_L = "FBE6C8", "B3781F"        # 설정에 따라

PANEL_F, PANEL_L = "FAFAFA", "999999"    # 결론 패널
INFO_F, INFO_L = "F3F6FA", "8CA0BE"      # 하단 정보 패널

KIND = {
    "self": (SELF_F, SELF_L),
    "store": (STORE_F, STORE_L),
    "api": (API_F, API_L),
    "opt": (OPT_F, OPT_L),
}

BADGE = {"yes": (STORE_L, "⏱"), "part": (OPT_L, "◐"), "no": (API_L, "✕")}

SLIDE_W, SLIDE_H = 13.33, 7.5
MARGIN = 0.42


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # Blank


def _style_runs(tf, size, bold, color, font=KR, space=None):
    for para in tf.paragraphs:
        if space is not None:
            para.line_spacing = space
        for run in para.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = _rgb(color)


def text(slide, x, y, w, h, content, *, size=11, bold=False, color=BODY,
         font=KR, align=PP_ALIGN.LEFT, wrap=False, space=None):
    """단순 텍스트 상자. content는 문자열 또는 줄 목록."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = content if isinstance(content, (list, tuple)) else [content]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.alignment = align
    _style_runs(tf, size, bold, color, font, space)
    return box


def rich(slide, x, y, w, h, parts, *, align=PP_ALIGN.LEFT, wrap=False):
    """한 줄 안에서 서식이 다른 조각들. parts = [(글자, dict(size,bold,color,font)), ...]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    for content, style in parts:
        run = para.add_run()
        run.text = content
        run.font.name = style.get("font", KR)
        run.font.size = Pt(style.get("size", 10))
        run.font.bold = style.get("bold", False)
        run.font.color.rgb = _rgb(style.get("color", BODY))
    return box


def header(slide, title, subtitle=None):
    """상단 네이비 바 + 제목(+ 바 안 부제)."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(0.82))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(NAVY)
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(slide, MARGIN, 0.13, 12.2, 0.34, title, size=21, bold=True, color=WHITE, wrap=True)
    if subtitle:
        text(slide, MARGIN, 0.49, 12.4, 0.24, subtitle, size=10, color=SUBTITLE, font=MONO)


def section(slide, x, y, label, *, w=7.0):
    """① 적재   POST /memories 같은 구역 제목."""
    return text(slide, x, y, w, 0.3, label, size=14, bold=True, color=NAVY, wrap=True)


def box(slide, x, y, w, h, lines, *, kind="self", size=10.5, bold=False,
        dashed=False, radius=0.12):
    """흐름도의 둥근사각형 박스."""
    fill, line = KIND[kind]
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.color.rgb = _rgb(line)
    shp.line.width = Pt(1.3)
    if dashed:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    seq = lines if isinstance(lines, (list, tuple)) else [lines]
    for i, line in enumerate(seq):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line if isinstance(line, str) else line[0]
        para.alignment = PP_ALIGN.CENTER
    # 첫 줄만 굵게 쓰는 원본 패턴 지원
    for i, para in enumerate(tf.paragraphs):
        want_bold = bold if i == 0 else False
        if not isinstance(seq[i], str):
            want_bold = seq[i][1]
        for run in para.runs:
            run.font.name = KR
            run.font.size = Pt(size)
            run.font.bold = want_bold
            run.font.color.rgb = _rgb(BODY)
    return shp


def badge(slide, cx, cy, kind, *, d=0.24, size=9):
    """박스 우상단에 붙는 계측 배지. cx, cy는 왼쪽 위 좌표."""
    color, glyph = BADGE[kind]
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.color.rgb = _rgb(WHITE)
    shp.line.width = Pt(1.2)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].text = glyph
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _style_runs(tf, size, True, WHITE)
    return shp


def arrow(slide, x, y, *, w=0.17, h=0.11):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.5
    shp.adjustments[1] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(GRAY)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def chip(slide, x, y, kind, *, w=0.2, h=0.15, radius=0.25):
    """범례/결론에 쓰는 작은 색 조각."""
    fill, line = KIND[kind]
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.color.rgb = _rgb(line)
    shp.line.width = Pt(1.0)
    shp.shadow.inherit = False
    return shp


def dot(slide, x, y, kind, *, d=0.12):
    """목록 왼쪽의 작은 분류 원."""
    fill, line = KIND[kind]
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.color.rgb = _rgb(line)
    shp.line.width = Pt(1.0)
    shp.shadow.inherit = False
    return shp


def pill(slide, x, y, w, h, label, *, color=STORE_L, size=12):
    """구역 머리말 알약(초록/빨강)."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.2
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    text(slide, x + 0.14, y + 0.07, w - 0.28, 0.24, label, size=size, bold=True, color=WHITE)
    return shp


def panel(slide, x, y, w, h, *, style="conclusion", radius=0.06):
    """하단 패널. style=conclusion(FAFAFA) 또는 info(F3F6FA)."""
    fill, line = (PANEL_F, PANEL_L) if style == "conclusion" else (INFO_F, INFO_L)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.color.rgb = _rgb(line)
    shp.line.width = Pt(1.1)
    shp.shadow.inherit = False
    return shp


def footer(slide, note, *, y=7.14):
    return text(slide, MARGIN, y, 12.5, 0.24, note, size=9.5, color=FOOT, wrap=True)


def legend(slide, y=1.0, items=(("self", "MemMachine 자체"), ("store", "외부 저장소"),
                                ("api", "외부 API"), ("opt", "설정에 따라"))):
    """상단 색 범례 한 줄. 원본 좌표 간격 그대로."""
    x = MARGIN
    for kind, label in items:
        chip(slide, x, y, kind, w=0.26, h=0.17)
        # 라벨 상자는 글자보다 넉넉히 잡는다. 왼쪽 정렬이라 시각 결과는 원본과 같고,
        # 폭 검사에서 가짜 경고만 사라진다.
        text(slide, x + 0.33, y, 1.35, 0.2, label, size=11)
        x += 1.55
    return x


def rect(slide, x, y, w, h, fill, line=None, lw=1.0):
    """각진 사각형 한 칸. 표를 직접 그릴 때 쓴다."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def arrow_down(slide, x, y, *, w=0.11, h=0.17):
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.5
    shp.adjustments[1] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(GRAY)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def lead_line(slide, x, y, w, lead, rest, *, size=11, h=0.24):
    """앞머리 어절만 굵은 결론 한 줄."""
    return rich(slide, x, y, w, h, [
        (lead, {"size": size, "bold": True}),
        (rest, {"size": size}),
    ])


def table(slide, x, y, w, col_w, rows, *, header=None, header_h=0.28, row_h=0.30,
          header_size=10, body_size=9, body_fill=SELF_F, row_fills=None,
          mono_cols=(), center_cols=(), emphasis_rows=(), emphasis_color=API_L,
          pad=0.07):
    """헤더행(네이비) + 데이터행을 도형으로 직접 그린다.

    col_w 합은 w와 같아야 한다. 어떤 행의 칸 수가 열 수보다 적으면
    마지막 칸이 남은 폭을 전부 차지한다(병합 효과).
    mono_cols/center_cols는 열 번호 집합, emphasis_rows는 행 번호 집합.
    반환값은 표의 아래쪽 y좌표.
    """
    top = y
    if header:
        rect(slide, x, top, w, header_h, NAVY)
        cx = x
        for i, cell in enumerate(header):
            align = PP_ALIGN.CENTER if i in center_cols else PP_ALIGN.LEFT
            text(slide, cx + pad, top + (header_h - 0.2) / 2, col_w[i] - 2 * pad, 0.2,
                 cell, size=header_size, bold=True, color=WHITE, align=align)
            cx += col_w[i]
        top += header_h

    for r, row in enumerate(rows):
        fill = row_fills[r] if row_fills else body_fill
        emphasized = r in emphasis_rows
        rect(slide, x, top, w, row_h, fill,
             emphasis_color if emphasized else WHITE, 1.5 if emphasized else 1.0)
        cx = x
        last = len(row) - 1
        for i, cell in enumerate(row):
            cw = sum(col_w[i:]) if i == last and last < len(col_w) - 1 else col_w[i]
            is_mono = i in mono_cols
            align = PP_ALIGN.CENTER if (i in center_cols or is_mono) else PP_ALIGN.LEFT
            text(slide, cx + pad, top + (row_h - 0.2) / 2, cw - 2 * pad, 0.2, cell,
                 size=(body_size + 2) if is_mono else body_size,
                 bold=is_mono, color=CODE if is_mono else BODY,
                 font=MONO if is_mono else KR, align=align)
            cx += cw
        top += row_h
    return top


def conclusion(slide, x, y, w, h, lines, *, title="결론", size=10.5, gap=0.27):
    """하단 결론 패널 + 앞머리 굵은 줄들."""
    panel(slide, x, y, w, h)
    ty = y + 0.13
    if title:
        text(slide, x + 0.3, ty, 3.0, 0.3, title, size=13, bold=True)
        ty += 0.36
    for lead, rest in lines:
        lead_line(slide, x + 0.3, ty, w - 0.6, lead, rest, size=size)
        ty += gap
    return ty
